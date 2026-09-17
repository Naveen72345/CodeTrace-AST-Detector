"""
app.py – Flask REST API server for CodeTrace.

Run:
    cd backend
    python app.py

Endpoints:
    POST   /api/auth/login
    GET    /api/auth/me
    POST   /api/auth/logout
    POST   /api/auth/update-profile
    POST   /api/auth/change-password
    POST   /api/analyze
    GET    /api/dashboard/stats
    GET    /api/results
    GET    /api/results/<report_id>
    GET    /api/history
    POST   /api/clear-history
    GET    /api/reports
    GET    /api/reports/export/<report_id>
    GET    /api/reports/export-all
    GET    /api/health
"""

from __future__ import annotations

import io
import uuid
from datetime import datetime, timezone
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_file, Response
from flask_cors import CORS

import database as db
import ast_engine as engine

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"

# ── App setup ─────────────────────────────────────────────────
app = Flask(
    __name__,
    static_folder=str(FRONTEND_DIR / "static"),
    template_folder=str(FRONTEND_DIR / "templates"),
)

# Allow requests from the frontend dev server AND direct file:// opens
CORS(
    app,
    origins=[
        "http://localhost:5500",
        "http://127.0.0.1:5500",
        "null",                    # file:// origin in some browsers
    ],
    supports_credentials=False,
)

# In-memory session store  {token: username}
# Tokens are random UUIDs; production systems should use JWT.
_sessions: dict[str, str] = {}

ALLOWED_EXT = {"py", "txt"}


# ─────────────────────────────────────────────────────────────
#  FRONTEND ROUTES
# ─────────────────────────────────────────────────────────────
@app.get("/")
def landing_page():
    return render_template("index.html")


@app.get("/<page>.html")
def frontend_page(page: str):
    allowed_pages = {
        "index",
        "dashboard",
        "analyze",
        "results",
        "history",
        "reports",
        "settings",
    }
    if page not in allowed_pages:
        return _err("Page not found.", 404)
    return render_template(f"{page}.html")


# ─────────────────────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────────────────────
def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _now_ms() -> int:
    return int(datetime.now(tz=timezone.utc).timestamp() * 1000)


def _format_date() -> str:
    return datetime.now(tz=timezone.utc).strftime("%b %d, %Y, %I:%M %p") + " UTC"


def _make_report_id() -> str:
    ts   = datetime.now(tz=timezone.utc).strftime("%y%m%d%H%M")
    rand = uuid.uuid4().hex[:4].upper()
    return f"CT-{ts}-{rand}"


def _classify(score: int) -> tuple[str, str]:
    """Returns (human label, key) for a similarity percentage."""
    if score >= 80:
        return "Very High Match / Duplicate", "duplicate"
    if score >= 70:
        return "High Match", "high"
    if score >= 50:
        return "Similar", "similar"
    return "Low Similarity", "low"


def _require_auth() -> str | None:
    """Extract + validate Bearer token. Returns username or None."""
    header = request.headers.get("Authorization", "")
    token  = header.removeprefix("Bearer ").strip()
    return _sessions.get(token)


def _err(msg: str, code: int = 400) -> tuple[Response, int]:
    return jsonify({"error": msg}), code


# ─────────────────────────────────────────────────────────────
#  AUTH ROUTES
# ─────────────────────────────────────────────────────────────

@app.post("/api/auth/login")
def login() -> tuple[Response, int]:
    data     = request.get_json(silent=True) or {}
    username = data.get("username", "").strip()
    password = data.get("password", "")

    if not username or not password:
        return _err("Username and password are required.")

    user = db.get_user_by_username(username)
    if not user or not db.verify_password(password, user["password_hash"]):
        return _err("Invalid username or password.", 401)

    token = str(uuid.uuid4())
    _sessions[token] = user["username"]

    return jsonify({
        "token": token,
        "user": {
            "username": user["username"],
            "email":    user["email"],
            "role":     user["role"],
        },
    }), 200


@app.get("/api/auth/me")
def get_me() -> tuple[Response, int]:
    username = _require_auth()
    if not username:
        return _err("Unauthorized.", 401)

    user = db.get_user_by_username(username)
    if not user:
        return _err("User not found.", 404)

    return jsonify({
        "username": user["username"],
        "email":    user["email"],
        "role":     user["role"],
    }), 200


@app.post("/api/auth/logout")
def logout() -> tuple[Response, int]:
    header = request.headers.get("Authorization", "")
    token  = header.removeprefix("Bearer ").strip()
    _sessions.pop(token, None)
    return jsonify({"ok": True}), 200


@app.post("/api/auth/update-profile")
def update_profile() -> tuple[Response, int]:
    username = _require_auth()
    if not username:
        return _err("Unauthorized.", 401)

    data         = request.get_json(silent=True) or {}
    new_username = data.get("username", "").strip()
    email        = data.get("email", "").strip()

    if not new_username:
        return _err("Username cannot be empty.")
    if not email or "@" not in email:
        return _err("A valid email address is required.")

    ok = db.update_user_profile(username, new_username, email)
    if not ok:
        return _err("That username is already taken.", 409)

    # Keep the session pointing to the new username
    header = request.headers.get("Authorization", "")
    token  = header.removeprefix("Bearer ").strip()
    if token in _sessions:
        _sessions[token] = new_username

    return jsonify({"ok": True, "username": new_username, "email": email}), 200


@app.post("/api/auth/change-password")
def change_password() -> tuple[Response, int]:
    username = _require_auth()
    if not username:
        return _err("Unauthorized.", 401)

    data        = request.get_json(silent=True) or {}
    current_pw  = data.get("currentPassword", "")
    new_pw      = data.get("newPassword", "")

    user = db.get_user_by_username(username)
    if not user or not db.verify_password(current_pw, user["password_hash"]):
        return _err("Current password is incorrect.")

    if len(new_pw) < 6:
        return _err("New password must be at least 6 characters long.")

    db.update_user_password(username, db.hash_password(new_pw))
    return jsonify({"ok": True}), 200


# ─────────────────────────────────────────────────────────────
#  ANALYSIS  ROUTE
# ─────────────────────────────────────────────────────────────

@app.post("/api/analyze")
def analyze() -> tuple[Response, int]:
    username = _require_auth()
    if not username:
        return _err("Unauthorized.", 401)

    uploaded = request.files.getlist("files")
    if not uploaded:
        return _err("No files received. Upload at least 2 files.")
    if len(uploaded) < 2:
        return _err("At least 2 files are required for comparison.")

    # ── Validate & read files ─────────────────────────────────
    file_data: list[dict] = []
    for f in uploaded:
        ext = _ext(f.filename)
        if ext not in ALLOWED_EXT:
            return _err(
                f'File "{f.filename}" is not supported. '
                "Only .py and .txt files are allowed."
            ), 415
        try:
            content = f.read().decode("utf-8", errors="replace")
        except Exception:
            return _err(f'Could not read "{f.filename}".')

        file_data.append({
            "name":    f.filename,
            "type":    ext,
            "size":    len(content.encode()),
            "content": content,
        })

    # ── Separate by type; create all same-type pairs ──────────
    py_files  = [fd for fd in file_data if fd["type"] == "py"]
    txt_files = [fd for fd in file_data if fd["type"] == "txt"]

    pairs: list[tuple[dict, dict]] = []
    for group in (py_files, txt_files):
        for i in range(len(group)):
            for j in range(i + 1, len(group)):
                pairs.append((group[i], group[j]))

    if not pairs:
        return _err(
            "No comparable file pairs found. "
            "Upload at least 2 .py files OR at least 2 .txt files."
        )

    # ── Run analysis and persist ──────────────────────────────
    db.record_files(file_data)

    now_ms = _now_ms()
    date   = _format_date()
    results: list[dict] = []

    for fa, fb in pairs:
        scores = engine.analyze_pair(fa["content"], fb["content"], fa["type"])
        label, key = _classify(scores["overall"])

        comp = {
            "reportId":         _make_report_id(),
            "file1":            fa["name"],
            "file2":            fb["name"],
            "fileType":         fa["type"],
            "overallSimilarity":scores["overall"],
            "codeStructure":    scores["codeStructure"],
            "astPattern":       scores["astPattern"],
            "logicSimilarity":  scores["logicSimilarity"],
            "textSimilarity":   scores["textSimilarity"],
            "classification":   label,
            "classificationKey":key,
            "detectionStatus":  label,
            "details":          scores["details"],
            "analysisDate":     date,
            "timestamp":        now_ms,
        }
        db.save_comparison(comp)
        results.append(comp)

    return jsonify({"results": results, "count": len(results)}), 200


# ─────────────────────────────────────────────────────────────
#  DASHBOARD
# ─────────────────────────────────────────────────────────────

@app.get("/api/dashboard/stats")
def dashboard_stats() -> tuple[Response, int]:
    username = _require_auth()
    if not username:
        return _err("Unauthorized.", 401)
    return jsonify(db.get_dashboard_stats()), 200


# ─────────────────────────────────────────────────────────────
#  RESULTS  /  HISTORY  /  REPORTS
# ─────────────────────────────────────────────────────────────

@app.get("/api/results")
def get_results() -> tuple[Response, int]:
    if not _require_auth():
        return _err("Unauthorized.", 401)
    return jsonify(db.get_all_comparisons()), 200


@app.get("/api/results/<report_id>")
def get_result(report_id: str) -> tuple[Response, int]:
    if not _require_auth():
        return _err("Unauthorized.", 401)
    comp = db.get_comparison_by_id(report_id)
    if not comp:
        return _err(f"Report '{report_id}' not found.", 404)
    return jsonify(comp), 200


@app.get("/api/history")
def get_history() -> tuple[Response, int]:
    if not _require_auth():
        return _err("Unauthorized.", 401)
    return jsonify(db.get_all_comparisons()), 200


@app.post("/api/clear-history")
def clear_history() -> tuple[Response, int]:
    if not _require_auth():
        return _err("Unauthorized.", 401)
    db.clear_history()
    return jsonify({"ok": True}), 200


@app.get("/api/reports")
def get_reports() -> tuple[Response, int]:
    if not _require_auth():
        return _err("Unauthorized.", 401)
    return jsonify(db.get_all_comparisons()), 200


# ─────────────────────────────────────────────────────────────
#  EXCEL EXPORT
# ─────────────────────────────────────────────────────────────

def _build_single_xlsx(comp: dict) -> io.BytesIO:
    """Generate a styled .xlsx for one report and return a BytesIO buffer."""
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    # ── Column widths ─────────────────────────────────────────
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 48

    # ── Style helpers ─────────────────────────────────────────
    hdr_font  = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill  = PatternFill(fgColor="2563EB", fill_type="solid")
    sub_fill  = PatternFill(fgColor="EFF6FF", fill_type="solid")
    lbl_font  = Font(bold=True, size=10)
    thin      = Side(style="thin")
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)
    center    = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left      = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    row = [0]   # mutable counter

    def next_row() -> int:
        row[0] += 1
        return row[0]

    def section_header(label: str) -> None:
        r = next_row()
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(row=r, column=1, value=label)
        c.font = hdr_font
        c.fill = hdr_fill
        c.alignment = center
        c.border = border

    def data_row(label: str, value: str | int) -> None:
        r = next_row()
        c1 = ws.cell(row=r, column=1, value=label)
        c2 = ws.cell(row=r, column=2, value=value)
        c1.font = lbl_font
        c1.fill = sub_fill
        c1.border = border
        c1.alignment = left
        c2.border = border
        c2.alignment = left

    def spacer() -> None:
        next_row()

    # ── Report Header ─────────────────────────────────────────
    section_header("CodeTrace — Analysis Report")
    data_row("Report ID",      comp["reportId"])
    data_row("File 1",         comp["file1"])
    data_row("File 2",         comp["file2"])
    data_row("File Type",      "Python" if comp["fileType"] == "py" else "Text")
    data_row("Analysis Date",  comp["analysisDate"])
    spacer()

    # ── Similarity Scores ─────────────────────────────────────
    section_header("Similarity Scores")
    data_row("Overall Similarity",  f"{comp['overallSimilarity']}%")
    data_row("Code Structure",       f"{comp['codeStructure']}%")
    data_row("AST Pattern",          f"{comp['astPattern']}%")
    data_row("Logic Similarity",     f"{comp['logicSimilarity']}%")
    data_row("Text Similarity",      f"{comp['textSimilarity']}%")
    spacer()

    # ── Classification ────────────────────────────────────────
    section_header("Classification")
    data_row("Classification",    comp["classification"])
    data_row("Detection Status",  comp["detectionStatus"])
    spacer()

    # ── Detection Details ─────────────────────────────────────
    if comp.get("details"):
        section_header("Detection Details")
        for detail in comp["details"]:
            data_row("•", detail)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _build_all_xlsx(comps: list[dict]) -> io.BytesIO:
    """Generate a multi-row .xlsx for all reports."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "All Reports"

    headers = [
        "Report ID", "File 1", "File 2", "File Type", "Analysis Date",
        "Overall (%)", "Code Structure (%)", "AST Pattern (%)",
        "Logic Similarity (%)", "Text Similarity (%)",
        "Classification", "Detection Status",
    ]

    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    hdr_fill = PatternFill(fgColor="2563EB", fill_type="solid")
    center   = Alignment(horizontal="center")

    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font      = hdr_font
        c.fill      = hdr_fill
        c.alignment = center
        letter = ws.cell(row=1, column=col).column_letter
        ws.column_dimensions[letter].width = 22

    for idx, comp in enumerate(comps, 2):
        row_data = [
            comp["reportId"], comp["file1"], comp["file2"],
            "Python" if comp["fileType"] == "py" else "Text",
            comp["analysisDate"],
            comp["overallSimilarity"], comp["codeStructure"],
            comp["astPattern"],        comp["logicSimilarity"],
            comp["textSimilarity"],    comp["classification"],
            comp["detectionStatus"],
        ]
        for col, val in enumerate(row_data, 1):
            ws.cell(row=idx, column=col, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@app.get("/api/reports/export/<report_id>")
def export_report(report_id: str):
    if not _require_auth():
        return _err("Unauthorized.", 401)

    comp = db.get_comparison_by_id(report_id)
    if not comp:
        return _err(f"Report '{report_id}' not found.", 404)

    try:
        buf = _build_single_xlsx(comp)
    except ImportError:
        return _err("openpyxl is not installed. Run: pip install openpyxl"), 500

    filename = f"CodeTrace_{report_id}.xlsx"
    return send_file(
        buf,
        mimetype=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=filename,
    )


@app.get("/api/reports/export-all")
def export_all_reports():
    if not _require_auth():
        return _err("Unauthorized.", 401)

    comps = db.get_all_comparisons()
    if not comps:
        return _err("No reports to export.", 404)

    try:
        buf = _build_all_xlsx(comps)
    except ImportError:
        return _err("openpyxl is not installed. Run: pip install openpyxl"), 500

    ts       = datetime.now(tz=timezone.utc).strftime("%Y%m%d%H%M")
    filename = f"CodeTrace_AllReports_{ts}.xlsx"
    return send_file(
        buf,
        mimetype=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=filename,
    )


# ─────────────────────────────────────────────────────────────
#  HEALTH CHECK
# ─────────────────────────────────────────────────────────────

@app.get("/api/health")
def health():
    return jsonify({"status": "ok", "service": "CodeTrace API v1.0"}), 200


# ─────────────────────────────────────────────────────────────
#  ENTRY POINT
# ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    db.init_db()
    print("=" * 54)
    print("  CodeTrace Backend  –  http://localhost:5000")
    print("  Default login: admin / admin123")
    print("=" * 54)
    app.run(debug=True, host="0.0.0.0", port=5000)
